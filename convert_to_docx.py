import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = create_element('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def generate_docx():
    doc = docx.Document()

    # Document Styles & Page Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("Capstone Enterprise Deliverables\nScalable Enterprise Architectural Deployments of Agentic AI Solutions")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(9, 13, 22)

    # Subtitle Metadata
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_meta = p_meta.add_run("Author: Manikandan  |  Roll Number: 2023103519  |  Directory: 2023103519-Manikandan  |  Date: October 2026")
    run_meta.font.name = "Arial"
    run_meta.font.size = Pt(10)
    run_meta.font.italic = True
    run_meta.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Read markdown source file
    with open('CAPSTONE_DELIVERABLES.md', 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code = False
    code_block = []

    for line in lines:
        raw_line = line.rstrip('\n')

        if raw_line.startswith('```'):
            if in_code:
                in_code = False
                p_code = doc.add_paragraph()
                p_code.paragraph_format.left_indent = Inches(0.3)
                p_code.paragraph_format.space_after = Pt(8)
                run = p_code.add_run('\n'.join(code_block))
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(30, 41, 59)
                code_block = []
            else:
                in_code = True
            continue

        if in_code:
            code_block.append(raw_line)
            continue

        if raw_line.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(raw_line[2:])
            run.font.name = 'Arial'
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(6, 182, 212)
        elif raw_line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(raw_line[3:])
            run.font.name = 'Arial'
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
        elif raw_line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(raw_line[4:])
            run.font.name = 'Arial'
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(139, 92, 246)
        elif raw_line.startswith('- ') or raw_line.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(raw_line[2:])
            run.font.name = 'Arial'
            run.font.size = Pt(10)
        elif raw_line.strip() == '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run('_________________________________________________________________________________')
            run.font.color.rgb = RGBColor(226, 232, 240)
        elif raw_line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(raw_line)
            run.font.name = 'Arial'
            run.font.size = Pt(10.5)

    doc.save('CAPSTONE_DELIVERABLES.docx')
    print('✅ CAPSTONE_DELIVERABLES.docx generated successfully!')

if __name__ == '__main__':
    generate_docx()
